"""
生成可执行的预算重分配 Excel 操作表
输入：budget_reallocation.json + attribution.json
输出：N-US_预算重分配执行操作表.xlsx
"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

attr = json.load(open('attribution.json', encoding='utf-8'))
br = json.load(open('budget_reallocation.json', encoding='utf-8'))
camps = attr['campaigns']

# ---------- 样式 ----------
INK = "24332B"; MUTED = "6B756E"; ACCENT = "4A7C59"; ACCENT2 = "C9713F"
DANGER = "B23A2E"; WARN = "C98A2D"; BG_HEAD = "24332B"; BG_ALT = "F6F3EC"
thin = Side(style='thin', color='E2DED3')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def set_cell(ws, r, c, v, bold=False, color=None, fill=None, align='left', num=None, font_size=11):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name='Microsoft YaHei', bold=bold, color=color or INK, size=font_size)
    cell.border = border
    cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    if fill:
        cell.fill = PatternFill('solid', fgColor=fill)
    if num is not None:
        cell.number_format = num
    return cell

def style_header(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        set_cell(ws, row, i, h, bold=True, color='FFFFFF', fill=BG_HEAD, align='center')

wb = Workbook()

# ================= Sheet 1 执行摘要 =================
ws = wb.active
ws.title = '执行摘要'
ws.sheet_view.showGridLines = False

cur, proj = br['current'], br['projected']
cuts, grows = br['cuts'], br['grows']

set_cell(ws, 1, 1, 'N-US 店铺近30天广告预算重分配执行操作表', bold=True, font_size=16, color=ACCENT)
set_cell(ws, 2, 1, '数据窗口：近30天（2026-07-16 ~ 08-15）｜基于领星ERP广告报表与MSKU利润报表自动生成', color=MUTED, font_size=10)

set_cell(ws, 4, 1, '一、当前投放基线', bold=True, font_size=13, color=ACCENT)
base = [
    ('广告总花费', cur['total_spend'], '$#,##0.00'),
    ('广告总销售额', cur['total_sales'], '$#,##0.00'),
    ('综合 ACOS', cur['acos'], '0.00"%"'),
    ('低效活动数', cuts['count'], '0'),
    ('高效活动数', grows['count'], '0'),
]
r = 5
for name, val, fmt in base:
    set_cell(ws, r, 1, name)
    set_cell(ws, r, 2, val, num=fmt, bold=True)
    r += 1

r += 1
set_cell(ws, r, 1, '二、预算重分配方案', bold=True, font_size=13, color=ACCENT)
r += 1
plan = [
    ('① 削减低效活动预算', f"对 {cuts['count']} 个低效活动削减 {cuts['ratio']*100:.0f}% 预算"),
    ('   释放预算', cuts['released_budget'], '$#,##0.00'),
    ('   预计销售额减少', cuts['sales_loss'], '$#,##0.00'),
    ('   预计利润损失（保守）', cuts['profit_loss'], '$#,##0.00'),
    ('② 加投高效活动预算', f"对 {grows['count']} 个高效活动按 ROAS 加权分配"),
    ('   加投预算', grows['add_budget'], '$#,##0.00'),
    ('   预计新增销售额', grows['add_sales'], '$#,##0.00'),
    ('   预计新增利润（保守）', grows['add_profit'], '$#,##0.00'),
]
for row in plan:
    if len(row) == 2:
        set_cell(ws, r, 1, row[0], bold=True)
        set_cell(ws, r, 2, row[1])
    else:
        set_cell(ws, r, 1, row[0])
        set_cell(ws, r, 2, row[1], num=row[2], bold=True, color=ACCENT)
    r += 1

r += 1
set_cell(ws, r, 1, '三、预期改善效果', bold=True, font_size=13, color=ACCENT)
r += 1
proj_rows = [
    ('优化后综合 ACOS', proj['acos'], '0.00"%"', f"改善 {(cur['acos']-proj['acos']):.2f}pp"),
    ('预期净利改善', proj['net_profit_improve'], '$#,##0.00', ''),
    ('净 ROI 改善率', proj['roi_improve_pct'], '0.0"%"', '相对当前广告花费'),
]
for name, val, fmt, note in proj_rows:
    set_cell(ws, r, 1, name)
    set_cell(ws, r, 2, val, num=fmt, bold=True, color=ACCENT)
    set_cell(ws, r, 3, note, color=MUTED, font_size=10)
    r += 1

r += 1
set_cell(ws, r, 1, '⚠ 执行注意', bold=True, font_size=12, color=DANGER)
notes = [
    '1) 高间接占比活动（自然单带动强）切成先降预算30%观察，勿一刀切暂停。',
    '2) 加投前确认高效活动对应SKU库存充足，避免断货导致排名下跌。',
    '3) 每周复核一次，ACOS改善至35%以下的低效活动可恢复预算。',
    '4) 本表为操作导引，具体订单/销售额以领星后台实时数据为准。',
]
for n in notes:
    set_cell(ws, r, 1, n, color=MUTED, font_size=10)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1

ws.column_dimensions['A'].width = 40
ws.column_dimensions['B'].width = 22
ws.column_dimensions['C'].width = 30
ws.column_dimensions['D'].width = 20

# ================= Sheet 2 削减操作清单 =================
ws2 = wb.create_sheet('削减操作清单')
cut_camps = sorted([c for c in camps if c['category'] == '低效' and c['spends'] > 50],
                   key=lambda x: x['spends'], reverse=True)
ratio = cuts['ratio']

set_cell(ws2, 1, 1, '低效活动削减操作清单（共 %d 个，削减预算 %d%%）' % (len(cut_camps), ratio*100), bold=True, font_size=14, color=DANGER)

headers = ['活动ID', '活动名称', '类型', '操作动作', '当前花费($)', '削减预算($)', '削减后预算($)',
           'ACOS(%)', 'ROAS', '净利润($)', '广告费/净利', '归因', '执行优先级']
widths = [16, 30, 8, 12, 12, 12, 13, 10, 8, 12, 12, 12, 10]
style_header(ws2, 2, headers, widths)

def action_for(c):
    # 净利为负 / 广告费严重超净利 → 暂停；否则 → 降预算
    gp = c.get('gross_profit')
    share = c.get('ad_share_of_profit')
    if gp is not None and gp <= 0:
        return '暂停', DANGER
    if share is not None and share >= 3:
        return '暂停', DANGER
    return '降预算', WARN

def priority_for(c):
    gp = c.get('gross_profit')
    share = c.get('ad_share_of_profit')
    acos = c.get('acos', 0)
    if gp is not None and gp <= 0 and acos >= 50:
        return 'P0 立即', DANGER
    if share is not None and share >= 3:
        return 'P1 本周', WARN
    return 'P2 次周', ACCENT

r = 3
for c in cut_camps:
    cut = c['spends'] * ratio
    new_spend = c['spends'] - cut
    action, acolor = action_for(c)
    prio, pcolor = priority_for(c)
    gp = c.get('gross_profit')
    share = c.get('ad_share_of_profit')
    vals = [
        c['campaign_id'], c['campaign_name'], c['sponsored_type'].upper(), action,
        round(c['spends'],2), round(cut,2), round(new_spend,2),
        round(c.get('acos',0),1), round(c.get('roas',0),2),
        (round(gp,2) if gp is not None else '—'),
        (round(share,2) if share is not None else '—'),
        c['category'], prio
    ]
    for i, v in enumerate(vals, 1):
        align = 'center' if i in (3,4,12,13) else 'right' if i in (5,6,7,8,9,10,11) else 'left'
        # 特殊填色
        if i == 4:
            set_cell(ws2, r, i, v, align=align, bold=True, color=acolor, fill=BG_ALT if r%2==0 else None)
        elif i == 13:
            set_cell(ws2, r, i, v, align=align, bold=True, color=pcolor, fill=BG_ALT if r%2==0 else None)
        elif i in (10,11):
            color = DANGER if (isinstance(v,(int,float)) and v < 0) else (ACCENT if v=='—' else INK)
            set_cell(ws2, r, i, v, align=align, color=color, num='0.00' if isinstance(v,(int,float)) else None, fill=BG_ALT if r%2==0 else None)
        else:
            set_cell(ws2, r, i, v, align=align, num='#,##0.00' if isinstance(v,(int,float)) and i in (5,6,7) else ('0.0' if i==8 else ('0.00' if isinstance(v,float) else None)), fill=BG_ALT if r%2==0 else None)
    r += 1

ws2.freeze_panes = 'A3'
ws2.auto_filter.ref = 'A2:M%d' % (r-1)

# ================= Sheet 3 加投操作清单 =================
ws3 = wb.create_sheet('加投操作清单')
grow_camps = sorted([c for c in camps if c['category'] == '高效' and c['spends'] > 20],
                    key=lambda x: x['roas'] or 0, reverse=True)[:20]
total_roas = sum(c['roas'] for c in grow_camps if c.get('roas'))
grow_budget = cuts['released_budget']

set_cell(ws3, 1, 1, '高效活动加投操作清单（共 %d 个，加投预算 $%.0f，按 ROAS 加权分配）' % (len(grow_camps), grow_budget), bold=True, font_size=14, color=ACCENT)

headers3 = ['活动ID', '活动名称', '类型', '当前花费($)', 'ROAS', 'ACOS(%)', '净利润($)',
            '加投预算($)', '加投后总预算($)', '预计新增销售($)', '执行优先级']
widths3 = [16, 30, 8, 12, 8, 10, 12, 12, 13, 14, 10]
style_header(ws3, 2, headers3, widths3)

r = 3
for c in grow_camps:
    add = c['roas'] / total_roas * grow_budget if c.get('roas') and total_roas else 0
    gp = c.get('gross_profit')
    vals = [
        c['campaign_id'], c['campaign_name'], c['sponsored_type'].upper(),
        round(c['spends'],2), round(c.get('roas',0),2), round(c.get('acos',0),1),
        (round(gp,2) if gp is not None else '—'),
        round(add,2), round(c['spends']+add,2), round(add*c.get('roas',0),2),
        'P1 本周' if c.get('roas',0) >= 4.5 else 'P2 次周'
    ]
    prio_color = ACCENT if c.get('roas',0) >= 4.5 else WARN
    for i, v in enumerate(vals, 1):
        align = 'center' if i in (3,11) else 'right' if i in (4,5,6,7,8,9,10) else 'left'
        if i == 8:
            set_cell(ws3, r, i, v, align=align, bold=True, num='#,##0.00', color=ACCENT, fill=BG_ALT if r%2==0 else None)
        elif i == 11:
            set_cell(ws3, r, i, v, align=align, bold=True, color=prio_color, fill=BG_ALT if r%2==0 else None)
        else:
            set_cell(ws3, r, i, v, align=align, num='#,##0.00' if i in (4,8,9,10) else ('0.00' if i==5 else ('0.0' if i==6 else None)), fill=BG_ALT if r%2==0 else None)
    r += 1

ws3.freeze_panes = 'A3'
ws3.auto_filter.ref = 'A2:K%d' % (r-1)

# ================= Sheet 4 执行步骤 =================
ws4 = wb.create_sheet('执行步骤')
set_cell(ws4, 1, 1, '分阶段落地执行步骤', bold=True, font_size=14, color=ACCENT)

steps = [
    ('阶段', '时间段', '操作内容', '验证标准'),
    ('第一步 立即止损', '当日', '暂停4个最严重亏损活动：328-N5691-SP-AUTO、328-N5691-SP-KW、168-14477-N-SP-AUTO、168-6484-N-SB-VIDEO', '当日广告花费下降约$1,555/30天'),
    ('第二步 批量降本', '本周', '对其余17个低效活动集体降预算50-60%；为20个高效活动提竞价与预算，优先ROAS>4.5x', '削减端释放约$3,256/30天'),
    ('第三步 持续优化', '每周', '复核ACOS与间接占比；加投后ACOS仍低的继续加预算；低效中ACOS改善至35%以下的可恢复', '综合ACOS向33.39%靠拢'),
    ('第四步 库存校验', '每次加投前', '确认高效活动对应SKU库存充足，避免断货', '库存可用天数≥30天'),
]
widths4 = [16, 10, 60, 30]
for i, (h, w) in enumerate(zip(steps[0], widths4), 1):
    ws4.column_dimensions[get_column_letter(i)].width = w
for rid, row in enumerate(steps, 1):
    for i, v in enumerate(row, 1):
        if rid == 1:
            set_cell(ws4, rid, i, v, bold=True, color='FFFFFF', fill=BG_HEAD, align='center')
        else:
            color = ACCENT if i == 1 else INK
            set_cell(ws4, rid, i, v, bold=(i==1), color=color, align='left' if i in (3,4) else 'center')

out = 'N-US_预算重分配执行操作表.xlsx'
wb.save(out)
print('已生成', out)